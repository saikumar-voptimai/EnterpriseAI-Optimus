"""Provider tests require no network; persistence tests use the PostgreSQL fixtures."""

import asyncio
import io
import json
import uuid
from datetime import timedelta

import httpx
import pytest
from sqlalchemy import func, select

from app.config import Settings
from app.models import (
    Conversation,
    Document,
    DocumentChunk,
    Incident,
    JobRun,
    Membership,
    Memory,
    Message,
    Note,
    Notification,
    Reminder,
    Report,
    ScheduledJob,
    Scope,
    ScopeGrant,
    User,
    Workspace,
)
from app.services.context import ContextService
from app.services.documents import extract_document
from app.services.errors import ProviderError, ServiceError
from app.services.gateway import Completion, ModelTurn, OpenRouterGateway
from app.services.jobs import JobService, utcnow
from app.worker import Worker


@pytest.fixture
def service_settings():
    return Settings(
        allow_external_ai=True,
        openrouter_api_key="test-key",
        openrouter_models="test/model",
        openrouter_default_model="test/model",
        max_context_chars=24000,
    )


def test_gateway_uses_fixed_url_allowlist_and_token_budget(service_settings):
    observed = []

    def handler(request):
        observed.append(request)
        return httpx.Response(
            200,
            json={
                "choices": [{"message": {"content": "A real response"}, "finish_reason": "stop"}],
                "usage": {
                    "prompt_tokens": 12,
                    "completion_tokens": 4,
                    "private_vendor_field": "excluded",
                },
            },
        )

    gateway = OpenRouterGateway(service_settings, httpx.MockTransport(handler))
    result = asyncio.run(gateway.complete([{"role": "user", "content": "hello"}]))
    assert result.content == "A real response"
    assert result.usage == {"prompt_tokens": 12, "completion_tokens": 4}
    assert str(observed[0].url) == "https://openrouter.ai/api/v1/chat/completions"
    assert json.loads(observed[0].content)["max_tokens"] == service_settings.max_output_tokens
    with pytest.raises(ServiceError, match="allowlist"):
        asyncio.run(
            gateway.complete(
                [{"role": "user", "content": "hello"}], model="https://attacker.invalid"
            )
        )
    assert len(observed) == 1


@pytest.mark.parametrize(
    "status,payload,retryable",
    [
        (500, {"error": "secret vendor body"}, True),
        (401, {"error": "secret vendor body"}, False),
        (200, {"choices": []}, False),
        (200, {"choices": [{"message": {"content": "partial"}, "finish_reason": "length"}]}, False),
        (200, {"choices": [{"message": {"content": "answer"}}], "usage": "invalid"}, False),
    ],
)
def test_gateway_never_turns_provider_failure_into_success(
    service_settings, status, payload, retryable
):
    gateway = OpenRouterGateway(
        service_settings, httpx.MockTransport(lambda _: httpx.Response(status, json=payload))
    )
    with pytest.raises(ProviderError) as caught:
        asyncio.run(gateway.complete([{"role": "user", "content": "private user input"}]))
    assert caught.value.retryable is retryable
    assert "secret vendor body" not in str(caught.value)
    assert "private user input" not in str(caught.value)


def test_gateway_timeout_is_explicit_and_disabled_ai_never_calls_transport(service_settings):
    def timeout(request):
        raise httpx.ReadTimeout("private body must not be echoed", request=request)

    gateway = OpenRouterGateway(service_settings, httpx.MockTransport(timeout))
    with pytest.raises(ProviderError, match="timed out") as caught:
        asyncio.run(gateway.complete([{"role": "user", "content": "hello"}]))
    assert caught.value.retryable
    service_settings.allow_external_ai = False
    with pytest.raises(ServiceError, match="disabled"):
        asyncio.run(gateway.complete([{"role": "user", "content": "hello"}]))


def test_supported_document_formats_are_parsed_and_invalid_uploads_rejected():
    from docx import Document as WordDocument
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

    for name, raw, expected in [
        ("note.txt", "héllo".encode(), "héllo"),
        ("note.md", b"# Heading", "# Heading"),
        ("data.csv", b'key,value\n"a,b",2\n', '"a,b",2'),
        ("data.json", b'{"value":42}', '"value": 42'),
    ]:
        parsed = extract_document(raw, name)
        assert expected in parsed.body
        assert "".join(parsed.chunks) == parsed.body
        assert len(parsed.content_hash) == 64
    stream = io.BytesIO()
    document = WordDocument()
    document.add_paragraph("Actual DOCX paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Actual table cell"
    document.save(stream)
    assert "Actual table cell" in extract_document(stream.getvalue(), "source.docx").body
    blank_pdf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(blank_pdf)
    text_pdf = io.BytesIO()
    text_writer = PdfWriter()
    page = text_writer.add_blank_page(width=200, height=200)
    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    page[NameObject("/Resources")] = DictionaryObject(
        {
            NameObject("/Font"): DictionaryObject(
                {
                    NameObject("/F1"): text_writer._add_object(font),
                }
            )
        }
    )
    content = DecodedStreamObject()
    content.set_data(b"BT /F1 12 Tf 10 50 Td (Actual PDF text) Tj ET")
    page[NameObject("/Contents")] = text_writer._add_object(content)
    text_writer.write(text_pdf)
    assert "Actual PDF text" in extract_document(text_pdf.getvalue(), "source.pdf").body
    for raw, filename in [
        (b"pretend", "source.exe"),
        (b"pretend", "source.pdf"),
        (b"{", "source.json"),
    ]:
        with pytest.raises(ServiceError):
            extract_document(raw, filename)
    with pytest.raises(ServiceError, match="No text"):
        extract_document(blank_pdf.getvalue(), "scan.pdf")
    with pytest.raises(ServiceError, match="size limit"):
        extract_document(b"12345", "note.txt", max_bytes=4)


def seed(session):
    owner = User(
        email=f"service-{uuid.uuid4()}@example.test",
        name="Owner",
        password_hash="x",
        active=True,
        clearance=3,
    )
    session.add(owner)
    session.flush()
    scope = Scope(name=f"Scope {uuid.uuid4()}", kind="enterprise")
    session.add(scope)
    session.flush()
    workspace = Workspace(
        name="Workspace",
        description="",
        scope_id=scope.id,
        created_by=owner.id,
        classification=2,
        external_ai_enabled=True,
    )
    session.add(workspace)
    session.flush()
    session.add(Membership(workspace_id=workspace.id, user_id=owner.id, role="manager"))
    note = Note(owner_id=owner.id, title="Private note", body="PRIVATE_ONLY_MARKER", kind="thought")
    session.add(note)
    session.flush()
    session.commit()
    return owner.id, workspace.id, note.id


class FakeGateway:
    def __init__(self, callback=None, fail=False):
        self.callback, self.fail, self.calls = callback, fail, []

    async def complete_turn(self, messages, model=None, request_id=None, **kwargs):
        self.calls.append(messages)
        if self.callback:
            self.callback()
        if self.fail:
            raise ProviderError(
                "Provider temporarily unavailable.", retryable=True, request_id=request_id
            )
        return ModelTurn("Persisted generated report", model, {"total_tokens": 10}, request_id, [])


def create_job(factory, settings, owner_id, **kwargs):
    with factory() as session:
        owner = session.get(User, owner_id)
        job = JobService(session, settings).create(
            owner,
            name=kwargs.pop("name", "Task"),
            instructions="Summarize the authorized sources",
            external_ai_enabled=True,
            next_run_at=kwargs.pop("next_run_at", utcnow() + timedelta(days=1)),
            **kwargs,
        )
        # These baseline worker checks cover migrated owner-bound schedules.
        # Manager approval and service principals have dedicated governance tests.
        if job.workspace_id:
            job.approval_status = "legacy"
            job.status = "active"
            job.execution_user_id = None
        session.commit()
        return job.id


def enqueue(factory, settings, owner_id, job_id, at=None):
    with factory() as session:
        run = JobService(session, settings).enqueue(
            session.get(User, owner_id), session.get(ScheduledJob, job_id), at
        )
        session.commit()
        return run.id


def test_context_shared_audience_never_contains_private_data(db, service_settings):
    owner_id, workspace_id, _ = seed(db)
    shared = Document(
        workspace_id=workspace_id,
        title="Shared",
        filename="a.txt",
        media_type="text/plain",
        content_hash="a" * 64,
        body="SHARED_MARKER",
        created_by=owner_id,
    )
    db.add(shared)
    db.flush()
    db.add(DocumentChunk(document_id=shared.id, ordinal=0, body=shared.body))
    public_conversation = Conversation(workspace_id=workspace_id, title="Shared conversation")
    private_conversation = Conversation(owner_id=owner_id, title="Private conversation")
    db.add_all([public_conversation, private_conversation])
    db.flush()
    owner = db.get(User, owner_id)
    service = ContextService(db, service_settings)
    shared_context = service.for_conversation(owner, public_conversation, "Summarize")
    private_context = service.for_conversation(owner, private_conversation, "Summarize")
    assert "PRIVATE_ONLY_MARKER" not in json.dumps(shared_context.messages)
    assert "SHARED_MARKER" in json.dumps(shared_context.messages)
    assert "PRIVATE_ONLY_MARKER" in json.dumps(private_context.messages)
    assert all(source["version"] for source in shared_context.sources + private_context.sources)
    assert (
        sum(len(message["content"]) for message in shared_context.messages)
        <= service_settings.max_context_chars
    )


def test_personal_summary_grants_do_not_override_workspace_external_ai_switch(db, service_settings):
    owner_id, workspace_id, _ = seed(db)
    viewer = User(
        email=f"reviewer-{uuid.uuid4()}@example.test",
        name="Reviewer",
        password_hash="x",
        active=True,
        clearance=3,
    )
    db.add(viewer)
    db.flush()
    workspace = db.get(Workspace, workspace_id)
    db.add(ScopeGrant(user_id=viewer.id, scope_id=workspace.scope_id, can_review=True))
    db.add(
        Incident(
            workspace_id=workspace_id,
            title="Publication",
            summary="SUMMARY_ALLOWED",
            decision="PUBLISHED_DECISION",
            severity="high",
            published=True,
            reported_by=owner_id,
            request_id=str(uuid.uuid4()),
        )
    )
    conversation = Conversation(owner_id=viewer.id, title="Personal")
    db.add(conversation)
    db.flush()
    service = ContextService(db, service_settings)
    allowed = json.dumps(service.for_conversation(viewer, conversation, "Summarize").messages)
    assert "SUMMARY_ALLOWED" in allowed
    assert "PUBLISHED_DECISION" in allowed
    workspace.external_ai_enabled = False
    db.flush()
    assert "SUMMARY_ALLOWED" not in json.dumps(
        service.for_conversation(viewer, conversation, "Summarize").messages
    )


def test_worker_occurrences_reports_and_due_reminders_are_persisted_once(
    db, db_factory, service_settings
):
    owner_id, _, note_id = seed(db)
    due = utcnow() - timedelta(minutes=1)
    db.add(
        Reminder(
            owner_id=owner_id, note_id=note_id, title="Due reminder", due_at=due, status="open"
        )
    )
    db.commit()
    job_id = create_job(db_factory, service_settings, owner_id, next_run_at=due)
    worker = Worker(db_factory, service_settings, FakeGateway())
    assert worker.schedule_due() == 1
    assert worker.schedule_due() == 0
    assert worker.notify_due_reminders() == 1
    assert worker.notify_due_reminders() == 0
    lease = worker.claim_one()
    assert lease
    asyncio.run(worker.execute(lease))
    asyncio.run(worker.execute(lease))
    with db_factory() as session:
        run = session.scalar(select(JobRun).where(JobRun.job_id == job_id))
        assert run.status == "succeeded" and run.result == "Persisted generated report"
        assert session.scalar(select(func.count()).select_from(Report)) == 1
        assert session.scalar(select(func.count()).select_from(Notification)) == 2
        assert session.get(ScheduledJob, job_id).status == "completed"


def test_worker_retries_transient_failure_then_persists_terminal_error(
    db, db_factory, service_settings
):
    owner_id, _, _ = seed(db)
    job_id = create_job(db_factory, service_settings, owner_id)
    run_id = enqueue(db_factory, service_settings, owner_id, job_id)
    worker = Worker(db_factory, service_settings, FakeGateway(fail=True))
    for attempt in range(1, 4):
        lease = worker.claim_one()
        assert lease
        asyncio.run(worker.execute(lease))
        with db_factory() as session:
            run = session.get(JobRun, run_id)
            assert run.attempts == attempt
            assert run.status == ("failed" if attempt == 3 else "queued")
            assert lease.request_id in run.error
            run.available_at = utcnow() - timedelta(seconds=1)
            session.commit()
    with db_factory() as session:
        assert session.scalar(select(func.count()).select_from(Report)) == 0
        assert session.scalar(select(func.count()).select_from(Notification)) == 1


def test_worker_fences_stale_lease_and_revocation_before_result_commit(
    db, db_factory, service_settings
):
    owner_id, workspace_id, _ = seed(db)
    job_id = create_job(db_factory, service_settings, owner_id, workspace_id=workspace_id)
    stale_run_id = enqueue(db_factory, service_settings, owner_id, job_id)

    def supersede():
        with db_factory() as session:
            run = session.get(JobRun, stale_run_id)
            run.lease_token = str(uuid.uuid4())
            session.commit()

    worker = Worker(db_factory, service_settings, FakeGateway(supersede))
    lease = worker.claim_one()
    asyncio.run(worker.execute(lease))
    with db_factory() as session:
        assert session.get(JobRun, stale_run_id).result is None
        assert session.scalar(select(func.count()).select_from(Report)) == 0

    revoked_run_id = enqueue(db_factory, service_settings, owner_id, job_id)

    def revoke():
        with db_factory() as session:
            member = session.get(Membership, (workspace_id, owner_id))
            member.role = "viewer"
            session.commit()

    worker.gateway = FakeGateway(revoke)
    asyncio.run(worker.execute(worker.claim_one()))
    with db_factory() as session:
        assert session.get(JobRun, revoked_run_id).status == "blocked"
        assert session.scalar(select(func.count()).select_from(Report)) == 0
        assert session.scalar(select(func.count()).select_from(Notification)) == 0


def test_manual_dependency_graph_uses_same_occurrence_and_rejects_private_to_shared(
    db, db_factory, service_settings
):
    owner_id, workspace_id, _ = seed(db)
    schedule = utcnow() + timedelta(days=1)
    dependency_id = create_job(
        db_factory, service_settings, owner_id, next_run_at=schedule, name="Dependency"
    )
    job_id = create_job(
        db_factory, service_settings, owner_id, next_run_at=schedule, dependency_ids=[dependency_id]
    )
    occurrence = utcnow()
    enqueue(db_factory, service_settings, owner_id, job_id, occurrence)
    enqueue(db_factory, service_settings, owner_id, job_id, occurrence)
    gateway = FakeGateway()
    worker = Worker(db_factory, service_settings, gateway)
    asyncio.run(worker.execute(worker.claim_one()))
    asyncio.run(worker.execute(worker.claim_one()))
    with db_factory() as session:
        runs = session.scalars(select(JobRun)).all()
        assert len(runs) == 2 and all(
            run.scheduled_for == occurrence and run.status == "succeeded" for run in runs
        )
        assert "Persisted generated report" in json.dumps(gateway.calls[1])
        with pytest.raises(ServiceError, match="accountable owner and audience"):
            JobService(session, service_settings).create(
                session.get(User, owner_id),
                name="Invalid shared task",
                instructions="Summarize",
                workspace_id=workspace_id,
                next_run_at=schedule,
                external_ai_enabled=True,
                dependency_ids=[dependency_id],
            )


@pytest.mark.parametrize("change", ["deactivate", "forget_memory"])
def test_worker_discards_output_when_owner_or_private_sources_change(
    db, db_factory, service_settings, change
):
    owner_id, _, note_id = seed(db)
    memory = Memory(owner_id=owner_id, note_id=note_id, text="Retained private memory")
    db.add(memory)
    db.commit()
    memory_id = memory.id
    job_id = create_job(db_factory, service_settings, owner_id)
    run_id = enqueue(db_factory, service_settings, owner_id, job_id)

    def change_access():
        with db_factory() as session:
            if change == "deactivate":
                session.get(User, owner_id).active = False
            else:
                session.get(Memory, memory_id).forgotten_at = utcnow()
            session.commit()

    worker = Worker(db_factory, service_settings, FakeGateway(change_access))
    asyncio.run(worker.execute(worker.claim_one()))
    with db_factory() as session:
        run = session.get(JobRun, run_id)
        assert run.status == "blocked" and run.result is None
        assert session.scalar(select(func.count()).select_from(Report)) == 0


@pytest.mark.parametrize("revoke", ["grant", "external_ai"])
def test_context_does_not_replay_history_or_reports_after_source_permission_revocation(
    db, service_settings, revoke
):
    owner_id, workspace_id, _ = seed(db)
    reviewer = User(
        email=f"reviewer-{uuid.uuid4()}@example.test",
        name="Reviewer",
        password_hash="x",
        active=True,
        clearance=3,
    )
    db.add(reviewer)
    db.flush()
    workspace = db.get(Workspace, workspace_id)
    grant = ScopeGrant(user_id=reviewer.id, scope_id=workspace.scope_id, can_review=True)
    incident = Incident(
        workspace_id=workspace_id,
        title="Publication",
        summary="ORIGINAL_PUBLISHED_SUMMARY",
        severity="high",
        published=True,
        reported_by=owner_id,
        request_id=str(uuid.uuid4()),
    )
    conversation = Conversation(owner_id=reviewer.id, title="Private history")
    db.add_all([grant, incident, conversation])
    db.flush()
    source = {
        "type": "incident",
        "id": incident.id,
        "title": incident.title,
        "version": incident.updated_at.isoformat(),
        "summary_only": True,
    }
    report = Report(
        owner_id=reviewer.id,
        title="Prior brief",
        body="REPORT_DERIVED_SECRET",
        source_refs=[source],
    )
    db.add(report)
    db.flush()
    now = utcnow()
    db.add_all(
        [
            Message(
                conversation_id=conversation.id,
                role="user",
                author_id=reviewer.id,
                content="Earlier request",
                created_at=now,
            ),
            Message(
                conversation_id=conversation.id,
                role="assistant",
                content="HISTORY_DERIVED_SECRET",
                created_at=now + timedelta(microseconds=1),
                source_refs=[
                    {
                        "type": "report",
                        "id": report.id,
                        "title": report.title,
                        "version": report.created_at.isoformat(),
                    }
                ],
            ),
        ]
    )
    db.flush()
    service = ContextService(db, service_settings)
    before = json.dumps(service.for_conversation(reviewer, conversation, "Continue").messages)
    assert "REPORT_DERIVED_SECRET" in before and "HISTORY_DERIVED_SECRET" in before
    if revoke == "grant":
        grant.can_review = False
    else:
        workspace.external_ai_enabled = False
    db.flush()
    after = json.dumps(service.for_conversation(reviewer, conversation, "Continue").messages)
    assert "REPORT_DERIVED_SECRET" not in after
    assert "HISTORY_DERIVED_SECRET" not in after
    assert "ORIGINAL_PUBLISHED_SUMMARY" not in after


def test_context_omits_historical_assistant_memory_after_forgetting(db, service_settings):
    owner_id, _, note_id = seed(db)
    owner = db.get(User, owner_id)
    memory = Memory(owner_id=owner_id, note_id=note_id, text="FORGOTTEN_MEMORY_MARKER")
    conversation = Conversation(owner_id=owner_id, title="Private history")
    db.add_all([memory, conversation])
    db.flush()
    now = utcnow()
    db.add_all(
        [
            Message(
                conversation_id=conversation.id,
                role="user",
                author_id=owner_id,
                content="Earlier request",
                created_at=now,
            ),
            Message(
                conversation_id=conversation.id,
                role="assistant",
                content="REPLAY_MEMORY_MARKER",
                created_at=now + timedelta(microseconds=1),
                source_refs=[
                    {
                        "type": "memory",
                        "id": memory.id,
                        "title": "Memory",
                        "version": memory.updated_at.isoformat(),
                    }
                ],
            ),
        ]
    )
    db.flush()
    service = ContextService(db, service_settings)
    assert "REPLAY_MEMORY_MARKER" in json.dumps(
        service.for_conversation(owner, conversation, "Continue").messages
    )
    memory.forgotten_at = utcnow()
    db.flush()
    after = json.dumps(service.for_conversation(owner, conversation, "Continue").messages)
    assert "REPLAY_MEMORY_MARKER" not in after and "FORGOTTEN_MEMORY_MARKER" not in after


def test_context_preserves_provenance_of_history_outside_current_retrieval_window(
    db, service_settings
):
    owner_id, _, note_id = seed(db)
    owner, note = db.get(User, owner_id), db.get(Note, note_id)
    conversation = Conversation(owner_id=owner_id, title="Private history")
    db.add(conversation)
    now = utcnow()
    for index in range(35):
        db.add(
            Note(
                owner_id=owner_id,
                title=f"Recent note {index}",
                body="Recent fact",
                kind="thought",
                created_at=now + timedelta(microseconds=index),
            )
        )
    db.flush()
    source = {
        "type": "note",
        "id": note.id,
        "title": note.title,
        "version": note.updated_at.isoformat(),
    }
    db.add_all(
        [
            Message(
                conversation_id=conversation.id,
                role="user",
                content="Earlier question",
                created_at=now,
            ),
            Message(
                conversation_id=conversation.id,
                role="assistant",
                content="OLDER_DERIVED_FACT",
                source_refs=[source],
                created_at=now + timedelta(seconds=1),
            ),
        ]
    )
    db.flush()
    service = ContextService(db, service_settings)
    bundle = service.for_conversation(owner, conversation, "Continue")
    assert "PRIVATE_ONLY_MARKER" not in bundle.messages[0]["content"]
    assert "OLDER_DERIVED_FACT" in json.dumps(bundle.messages)
    assert note_id in {reference["id"] for reference in bundle.sources}
    db.add_all(
        [
            Message(
                conversation_id=conversation.id,
                role="user",
                content="Continue",
                created_at=now + timedelta(seconds=2),
            ),
            Message(
                conversation_id=conversation.id,
                role="assistant",
                content="NEWER_DERIVED_FACT",
                source_refs=bundle.sources,
                created_at=now + timedelta(seconds=3),
            ),
        ]
    )
    db.delete(note)
    db.flush()
    after = json.dumps(service.for_conversation(owner, conversation, "Continue again").messages)
    assert "OLDER_DERIVED_FACT" not in after and "NEWER_DERIVED_FACT" not in after
