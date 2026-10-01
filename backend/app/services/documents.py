"""Text extraction from a small explicit format allowlist; no arbitrary file paths."""

import csv
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import PurePath

from app.config import get_settings
from app.services.errors import ServiceError

MAX_TEXT_CHARS = 1_000_000
CHUNK_CHARS = 1800
MAX_CHUNKS = 600
MAX_ZIP_EXPANDED_BYTES = 20_971_520
MEDIA_TYPES = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".csv": "text/csv",
    ".json": "application/json",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


@dataclass(frozen=True)
class ExtractedDocument:
    body: str
    content_hash: str
    media_type: str
    chunks: list[str]


def _utf8(data: bytes) -> str:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise ServiceError("Text files must use UTF-8 encoding.", 415) from exc
    if "\x00" in text:
        raise ServiceError("Binary content is not accepted as a text document.", 415)
    return text


def extract_document(
    data: bytes, filename: str, media_type: str | None = None, max_bytes: int | None = None
) -> ExtractedDocument:
    maximum = max_bytes if max_bytes is not None else get_settings().max_upload_bytes
    if not data:
        raise ServiceError("The uploaded file is empty.")
    if len(data) > maximum:
        raise ServiceError("The uploaded file exceeds the configured size limit.", 413)
    suffix = PurePath(filename.lower()).suffix
    if suffix not in MEDIA_TYPES:
        raise ServiceError("Supported files are TXT, MD, CSV, JSON, PDF, and DOCX.", 415)
    # A supplied MIME header is not trusted to choose a parser.
    try:
        if suffix in {".txt", ".md"}:
            text = _utf8(data)
        elif suffix == ".csv":
            output = io.StringIO()
            writer = csv.writer(output, lineterminator="\n")
            for row in csv.reader(io.StringIO(_utf8(data)), strict=True):
                writer.writerow(row)
                if output.tell() > MAX_TEXT_CHARS:
                    raise ServiceError("Extracted document text exceeds the size limit.", 413)
            text = output.getvalue()
        elif suffix == ".json":
            text = json.dumps(json.loads(_utf8(data)), ensure_ascii=False, indent=2)
        elif suffix == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise ServiceError("The file is not a valid PDF.", 415)
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ServiceError("Encrypted PDFs are not supported.", 415)
            if len(reader.pages) > 300:
                raise ServiceError("PDFs are limited to 300 pages.", 413)
            pages: list[str] = []
            total = 0
            for page in reader.pages:
                # Reject oversized decoded content before invoking text extraction.
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > MAX_ZIP_EXPANDED_BYTES:
                    raise ServiceError("PDF page content exceeds the extraction limit.", 413)
                page_text = page.extract_text() or ""
                total += len(page_text)
                if total > MAX_TEXT_CHARS:
                    raise ServiceError("Extracted document text exceeds the size limit.", 413)
                pages.append(page_text)
            text = "\n\n".join(pages)
        else:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                members = archive.infolist()
                if (
                    len(members) > 1000
                    or sum(item.file_size for item in members) > MAX_ZIP_EXPANDED_BYTES
                ):
                    raise ServiceError("The DOCX archive exceeds the extraction limit.", 413)
                if "word/document.xml" not in archive.namelist():
                    raise ServiceError("The file is not a valid DOCX.", 415)
                if any(item.flag_bits & 1 for item in members):
                    raise ServiceError("Encrypted DOCX files are not supported.", 415)
            from docx import Document

            document = Document(io.BytesIO(data))
            paragraphs = [p.text for p in document.paragraphs]
            for table in document.tables:
                paragraphs.extend("\t".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(paragraphs)
    except ServiceError:
        raise
    except ImportError as exc:
        raise ServiceError("The document parser is not installed on this server.", 503) from exc
    except Exception as exc:
        raise ServiceError("The document could not be parsed as the selected format.", 422) from exc
    text = text.replace("\x00", "").strip()
    if not text:
        raise ServiceError("No text was found. Scanned images require OCR before upload.", 422)
    if len(text) > MAX_TEXT_CHARS:
        raise ServiceError("Extracted document text exceeds the size limit.", 413)
    chunks = [text[offset : offset + CHUNK_CHARS] for offset in range(0, len(text), CHUNK_CHARS)]
    if len(chunks) > MAX_CHUNKS:
        raise ServiceError("The document exceeds the chunk limit.", 413)
    return ExtractedDocument(text, hashlib.sha256(data).hexdigest(), MEDIA_TYPES[suffix], chunks)


@dataclass(frozen=True)
class ChunkFragment:
    body: str
    start_offset: int
    end_offset: int
    location: dict


def chunk_text(
    text: str, target_tokens: int = 500, overlap_tokens: int = 75
) -> list[ChunkFragment]:
    """Token-bounded overlapping chunks with exact offsets and nearby headings.

    Prefer paragraph boundaries in the last quarter of each window. Offsets
    point into the revision's immutable extracted text, not arbitrary filenames.
    """
    import bisect
    import re
    import tiktoken

    if target_tokens < 32 or not 0 <= overlap_tokens < target_tokens:
        raise ValueError("Invalid chunk token limits")
    encoding = tiktoken.get_encoding("cl100k_base")
    tokens = encoding.encode(text, disallowed_special=())
    _, offsets = encoding.decode_with_offsets(tokens)
    offsets.append(len(text))
    heading_positions = [
        (match.start(), match.group(1).strip())
        for match in re.finditer(r"(?m)^#{1,6}\s+(.+)$", text)
    ]
    fragments, start = [], 0
    while start < len(tokens):
        end = min(start + target_tokens, len(tokens))
        char_start, char_end = offsets[start], offsets[end]
        if end < len(tokens):
            paragraph = text.rfind("\n\n", offsets[start + target_tokens * 3 // 4], char_end)
            if paragraph > char_start:
                end = max(start + 1, bisect.bisect_left(offsets, paragraph + 2))
                char_end = offsets[end]
        heading = next(
            (title for position, title in reversed(heading_positions) if position <= char_start),
            None,
        )
        location = {"start_offset": char_start, "end_offset": char_end}
        if heading:
            location["heading"] = heading
        fragments.append(ChunkFragment(text[char_start:char_end], char_start, char_end, location))
        if end == len(tokens):
            break
        start = max(start + 1, end - overlap_tokens)
    if len(fragments) > MAX_CHUNKS:
        raise ServiceError("The document exceeds the chunk limit.", 413)
    return fragments
