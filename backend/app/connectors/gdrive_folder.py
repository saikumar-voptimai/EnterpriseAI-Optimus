"""A Google Drive folder shared, read-only, with a Google service account.

The folder is shared with the service account's email as Viewer; access then
belongs to the workspace rather than to one person's Google sign-in. Only files
inside that folder (or its subfolders) are listed or read. Spreadsheets
(Google Sheets, Excel, CSV) are returned as bounded tables.
"""

import base64
import csv
import datetime as dt
import io
import json
import re
import time
from urllib.parse import quote
from cryptography.hazmat.primitives import hashes, serialization
from cryptography.hazmat.primitives.asymmetric import padding
from app.connectors.base import ConnectorHTTP
from app.services.errors import ProviderError, ServiceError

TOKEN_URL = "https://oauth2.googleapis.com/token"
DRIVE = "https://www.googleapis.com/drive/v3"
SCOPE = "https://www.googleapis.com/auth/drive.readonly"
FOLDER = "application/vnd.google-apps.folder"
SHEET = "application/vnd.google-apps.spreadsheet"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EXCEL_TYPES = {XLSX, "application/vnd.ms-excel.sheet.macroEnabled.12"}
CSV_TYPES = {"text/csv", "text/plain", "application/csv"}
MAX_BYTES = 20_000_000
MAX_ROWS, MAX_COLUMNS = 500, 60


def folder_id_from(value):
    """Accept a folder link or a bare folder ID."""
    value = (value or "").strip()
    match = re.search(r"/folders/([A-Za-z0-9_-]{10,})", value) or re.fullmatch(
        r"([A-Za-z0-9_-]{10,})", value
    )
    if not match:
        raise ServiceError("Paste the Google Drive folder link.", 422)
    return match.group(1)


def service_account(raw):
    try:
        info = json.loads(raw) if isinstance(raw, str) else dict(raw)
    except (ValueError, TypeError) as exc:
        raise ServiceError("Paste the service account key file contents (JSON).", 422) from exc
    if (
        info.get("type") != "service_account"
        or not str(info.get("client_email", "")).endswith(".iam.gserviceaccount.com")
        or "PRIVATE KEY" not in str(info.get("private_key", ""))
        or info.get("token_uri", TOKEN_URL) != TOKEN_URL
    ):
        raise ServiceError("This is not a Google service account key file.", 422)
    return {k: info[k] for k in ("client_email", "private_key", "private_key_id") if k in info}


def b64(data: bytes) -> str:
    return base64.urlsafe_b64encode(data).rstrip(b"=").decode()


def cell(value):
    if value is None or isinstance(value, (bool, int, str)):
        return value
    if isinstance(value, float):
        return value if value == value and abs(value) != float("inf") else None
    if isinstance(value, (dt.datetime, dt.date, dt.time)):
        return value.isoformat()
    return str(value)[:1000]


class DriveFolder:
    def __init__(self, config, credentials, http=None):
        self.folder_id = folder_id_from(config.get("folder_id", ""))
        self.account = service_account(credentials.get("service_account", {}))
        self.http = http or ConnectorHTTP(timeout=60)
        self._token = None

    def assertion(self, now=None):
        now = int(now or time.time())
        header = {"alg": "RS256", "typ": "JWT"}
        if self.account.get("private_key_id"):
            header["kid"] = self.account["private_key_id"]
        claims = {
            "iss": self.account["client_email"],
            "scope": SCOPE,
            "aud": TOKEN_URL,
            "iat": now,
            "exp": now + 3600,
        }
        signing_input = (
            b64(json.dumps(header, separators=(",", ":")).encode())
            + "."
            + b64(json.dumps(claims, separators=(",", ":")).encode())
        )
        try:
            key = serialization.load_pem_private_key(
                self.account["private_key"].encode(), password=None
            )
        except (ValueError, TypeError) as exc:
            raise ServiceError("The service account private key could not be read.", 422) from exc
        signature = key.sign(signing_input.encode(), padding.PKCS1v15(), hashes.SHA256())
        return signing_input + "." + b64(signature)

    async def token(self):
        if self._token:
            return self._token
        result = await self.http.json(
            "POST",
            TOKEN_URL,
            data={
                "grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer",
                "assertion": self.assertion(),
            },
        )
        if not result.get("access_token"):
            raise ProviderError("Google did not issue a service account token.")
        self._token = result["access_token"]
        return self._token

    async def _get(self, path, **params):
        return await self.http.json(
            "GET",
            DRIVE + path,
            headers={"Authorization": "Bearer " + await self.token()},
            params={"supportsAllDrives": "true", **params},
        )

    async def folder(self):
        meta = await self._get(
            "/files/" + quote(self.folder_id, safe=""), fields="id,name,mimeType"
        )
        if meta.get("mimeType") != FOLDER:
            raise ServiceError("The link does not point to a folder.", 422)
        return {"id": meta["id"], "name": meta.get("name", "Shared folder")}

    async def files(self, search="", limit=200):
        """Files in the folder and up to three levels of subfolders, newest first."""
        found, queue, depth = [], [(self.folder_id, "")], 0
        while queue and depth < 4 and len(found) < limit:
            next_queue = []
            for parent, prefix in queue:
                page = await self._get(
                    "/files",
                    q=f"'{parent}' in parents and trashed = false",
                    fields="files(id,name,mimeType,modifiedTime,size)",
                    pageSize=200,
                    orderBy="modifiedTime desc",
                    includeItemsFromAllDrives="true",
                )
                for item in page.get("files", []):
                    path = f"{prefix}{item.get('name', '')}"
                    if item.get("mimeType") == FOLDER:
                        next_queue.append((item["id"], path + "/"))
                    elif not search or search.lower() in path.lower():
                        found.append(
                            {
                                "id": item["id"],
                                "name": path,
                                "kind": (
                                    "spreadsheet"
                                    if item.get("mimeType") in {SHEET, *EXCEL_TYPES, *CSV_TYPES}
                                    else "file"
                                ),
                                "modified_at": item.get("modifiedTime"),
                            }
                        )
            queue, depth = next_queue, depth + 1
        found.sort(key=lambda f: f["modified_at"] or "", reverse=True)
        return found[:limit]

    async def _inside(self, file_id):
        """The file must sit within the shared folder; walk parents up to six levels."""
        current, meta = file_id, None
        for level in range(7):
            item = await self._get(
                "/files/" + quote(current, safe=""),
                fields="id,name,mimeType,parents,size,modifiedTime",
            )
            if level == 0:
                meta = item
            parents = item.get("parents") or []
            if self.folder_id in parents:
                return meta
            if not parents:
                break
            current = parents[0]
        raise ServiceError("This file is not in the shared folder.", 403)

    async def read_table(self, file_id, sheet=None, offset=0, limit=100):
        meta = await self._inside(file_id)
        mime = meta.get("mimeType", "")
        headers = {"Authorization": "Bearer " + await self.token()}
        if mime == SHEET:
            url, params = DRIVE + "/files/" + quote(file_id, safe="") + "/export", {
                "mimeType": XLSX
            }
        elif (
            mime in EXCEL_TYPES
            or mime in CSV_TYPES
            or meta.get("name", "").lower().endswith((".xlsx", ".xlsm", ".csv"))
        ):
            url, params = DRIVE + "/files/" + quote(file_id, safe=""), {
                "alt": "media",
                "supportsAllDrives": "true",
            }
        else:
            raise ServiceError(
                "Only spreadsheets (Google Sheets, Excel, CSV) can be read as tables.", 415
            )
        content = await self.http.request(
            "GET", url, headers=headers, params=params, max_bytes=MAX_BYTES
        )
        offset, limit = max(0, int(offset)), max(1, min(int(limit), MAX_ROWS))
        if mime in CSV_TYPES or meta.get("name", "").lower().endswith(".csv"):
            sheets = {
                "Sheet1": list(csv.reader(io.StringIO(content.decode("utf-8-sig", "replace"))))
            }
            chosen = "Sheet1"
        else:
            sheets, chosen = self._workbook(content, sheet)
        rows = [r for r in sheets[chosen] if any(v not in (None, "") for v in r)]
        header = [
            str(v) if v not in (None, "") else f"Column {i + 1}"
            for i, v in enumerate((rows[0] if rows else [])[:MAX_COLUMNS])
        ]
        body = rows[1:]
        window = body[offset : offset + limit]
        return {
            "file": meta.get("name"),
            "modified_at": meta.get("modifiedTime"),
            "sheets": list(sheets),
            "sheet": chosen,
            "columns": header,
            "rows": [dict(zip(header, (cell(v) for v in row[:MAX_COLUMNS]))) for row in window],
            "offset": offset,
            "total_rows": len(body),
            "truncated": offset + limit < len(body),
        }

    @staticmethod
    def _workbook(content, sheet):
        from openpyxl import load_workbook

        try:
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as exc:  # openpyxl raises many types for malformed files
            raise ServiceError("The spreadsheet could not be opened.", 422) from exc
        try:
            names = book.sheetnames
            chosen = sheet if sheet in names else names[0]
            values = {name: [] for name in names}
            for row in book[chosen].iter_rows(values_only=True, max_row=20_000):
                values[chosen].append(list(row))
            return values, chosen
        finally:
            book.close()
